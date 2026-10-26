from __future__ import annotations

from pathlib import Path

from docx import Document

ROOT = Path(__file__).parents[1]
REQUIRED = (
    ROOT / "README.md",
    ROOT / "docs" / "architecture.md",
    ROOT / "docs" / "security.md",
    ROOT / "docs" / "evaluation-report.md",
    ROOT / "docs" / "functional-acceptance-report.md",
    ROOT / "docs" / "project-documentation.md",
    ROOT / "docs" / "video-demo-script.md",
    ROOT / "docs" / "Week3-CodeReviewAgent-Project-Documentation.docx",
)
MARKDOWN_DOCUMENTS = REQUIRED[:-1]
STATUS_DOCUMENTS = (
    ROOT / "README.md",
    ROOT / "docs" / "evaluation-report.md",
    ROOT / "docs" / "functional-acceptance-report.md",
    ROOT / "docs" / "project-documentation.md",
    ROOT / "docs" / "video-demo-script.md",
)
PR_URLS = (
    "https://github.com/melepatt/code-review-agent-demo/pull/1",
    "https://github.com/melepatt/code-review-agent-demo/pull/2",
    "https://github.com/melepatt/code-review-agent-demo/pull/3",
    "https://github.com/melepatt/code-review-agent-demo/pull/6",
    "https://github.com/melepatt/code-review-agent-demo/pull/4",
    "https://github.com/melepatt/code-review-agent-demo/pull/5",
    "https://github.com/melepatt/code-review-agent-demo/pull/7",
    "https://github.com/melepatt/code-review-agent-demo/pull/9",
    "https://github.com/melepatt/code-review-agent-demo/pull/8",
    "https://github.com/melepatt/code-review-agent-demo/pull/10",
)
EVALUATION_ARTIFACTS = (
    "evaluation/runs/evaluation-20261026T002439043692Z-clean-controls.json",
    "evaluation/runs/evaluation-20261026T004421625474Z-seeded.json",
    "evaluation/runs/evaluation-20261026T004440793716Z-recovery.json",
)
FALSE_BASELINE_CLAIMS = (
    "live model evaluation: not run",
    "paid model evaluation and github review publication have not run",
    "live model evaluation: pass",
    "live model evaluation passed",
    "live accuracy passed",
    "ten-case sequence passed",
    "complete live catalog pass",
    "github review publication: pass",
    "github review publication passed",
)


def _normalize(text: str) -> str:
    return " ".join(text.casefold().replace("|", " ").replace("`", "").split())


def _assert_live_baseline_contract(text: str, source: str) -> None:
    normalized = _normalize(text)
    required = (
        "clean-controls-produced-findings",
        "10 percent",
        "14.29 percent",
        "93.33 percent",
        "one matched finding",
        "week 4",
    )
    alternatives = (
        ("0 percent duplicates", "duplicate rate 0 percent", "duplicate rate is 0 percent"),
        (
            "failed clean behavior",
            "failed clean-control behavior",
            "clean behavior failed",
            "clean-control behavior failed",
        ),
        (
            "failed recovery",
            "partial recovery failed",
            "failed malformed-response recovery",
            "malformed-response recovery failed",
        ),
        (
            "github review publication: not run",
            "github review publication was not run",
        ),
    )

    for phrase in required:
        assert phrase in normalized, f"{source} is missing {phrase!r}"
    for choices in alternatives:
        assert any(choice in normalized for choice in choices), (
            f"{source} is missing one of {choices!r}"
        )
    for false_claim in FALSE_BASELINE_CLAIMS:
        assert false_claim not in normalized, f"{source} contains stale claim {false_claim!r}"


def test_required_documentation_artifacts_exist_and_are_nonempty() -> None:
    missing = [str(path.relative_to(ROOT)) for path in REQUIRED if not path.is_file()]
    empty = [
        str(path.relative_to(ROOT))
        for path in REQUIRED
        if path.is_file() and path.stat().st_size == 0
    ]

    assert missing == []
    assert empty == []


def test_project_docx_is_valid_and_has_navigation_headings() -> None:
    document = Document(REQUIRED[-1])
    title = next(
        paragraph.text for paragraph in document.paragraphs if paragraph.style.name == "Title"
    )
    headings = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading")
    ]

    assert title == "Code Review Agent Project Documentation"
    assert len(headings) >= 6


def test_documentation_records_live_baseline_without_overstating_the_gate() -> None:
    text = "\n".join(path.read_text(encoding="utf-8") for path in MARKDOWN_DOCUMENTS)
    normalized = _normalize(text)

    assert "ten synthetic" in normalized
    assert "seven seeded" in normalized
    assert "two clean" in normalized
    assert "one controlled malformed" in normalized
    assert "clean controls" in normalized
    assert "seeded defects" in normalized
    assert "malformed-recovery" in normalized
    assert "at most once" in normalized
    for path in STATUS_DOCUMENTS:
        _assert_live_baseline_contract(
            path.read_text(encoding="utf-8"),
            str(path.relative_to(ROOT)),
        )
    for url in PR_URLS:
        assert url in text
    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    for artifact in EVALUATION_ARTIFACTS:
        assert f"]({artifact})" in readme


def test_project_docx_records_live_baseline_and_week_four_handoff() -> None:
    document = Document(REQUIRED[-1])
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for table in document.tables:
        text += "\n" + "\n".join(cell.text for row in table.rows for cell in row.cells)
    normalized = _normalize(text)

    assert "ten synthetic" in normalized
    assert "seven seeded" in normalized
    assert "two clean" in normalized
    assert "one controlled malformed" in normalized
    _assert_live_baseline_contract(text, "project DOCX")
