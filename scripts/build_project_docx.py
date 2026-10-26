from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F6F7"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_fill(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def set_cell_margins(
    cell, top: int = 110, start: int = 120, bottom: int = 110, end: int = 120
) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), BORDER)


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    properties.append(cant_split)


def set_keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_bullet(document: Document, text: str, level: int = 0) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.28 + (0.22 * level))
    paragraph.paragraph_format.first_line_indent = Inches(-0.2)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run("•  ")
    paragraph.add_run(text)


def add_number(document: Document, number: int, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.34)
    paragraph.paragraph_format.first_line_indent = Inches(-0.28)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(f"{number}.  ")
    paragraph.add_run(text)


def add_table(
    document: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    *,
    font_size: float = 9,
    vertical_margin: int = 110,
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for index, text in enumerate(headers):
        cell = header.cells[index]
        cell.width = Inches(widths[index])
        set_cell_fill(cell, NAVY)
        set_cell_margins(cell, top=vertical_margin, bottom=vertical_margin)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(font_size)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        cells = row.cells
        for index, text in enumerate(values):
            cell = cells[index]
            cell.width = Inches(widths[index])
            set_cell_margins(cell, top=vertical_margin, bottom=vertical_margin)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_fill(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            )
            run = paragraph.add_run(text)
            run.font.size = Pt(font_size)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run("Page ")
    run.font.size = Pt(9)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)


def configure(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for style_name, size, before, after in (
        ("Title", 25, 0, 14),
        ("Subtitle", 12, 0, 10),
        ("Heading 1", 17, 16, 7),
        ("Heading 2", 12.5, 11, 5),
        ("Heading 3", 11, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Normal" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = paragraph.add_run("Code Review Agent  |  ")
    label.font.size = Pt(9)
    label.font.color.rgb = RGBColor(89, 89, 89)
    add_page_number(paragraph)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    set_keep_with_next(paragraph)


def build(output: Path) -> None:
    document = Document()
    configure(document)

    cover = document.add_paragraph(style="Title")
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(118)
    cover.add_run("Code Review Agent Project Documentation")
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Week 3 Project Option 3E")
    detail = document.add_paragraph()
    detail.alignment = WD_ALIGN_PARAGRAPH.CENTER
    detail.paragraph_format.space_before = Pt(22)
    detail.add_run(
        "A production-minded Streamlit and LangGraph application for bounded, grounded, "
        "human-approved GitHub pull-request review"
    )
    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.space_before = Pt(30)
    date.add_run("25 October 2026")
    document.add_page_break()

    add_heading(document, "Document purpose")
    document.add_paragraph(
        "This report explains what the project does, why its safety boundaries matter, how the "
        "implementation works, and what evidence exists. The offline application and safeguards "
        "are implemented, and ten synthetic GitHub fixtures are pinned in the golden catalog. "
        "The live evaluation established a failed clean-control gate and a diagnostic ten-case "
        "baseline for Week 4. GitHub review publication has not run."
    )
    add_heading(document, "Contents")
    for item in (
        "Project overview",
        "Agent workflow",
        "System architecture",
        "Security and human control",
        "Testing and evaluation",
        "Development iterations",
        "Operation",
        "Results and limitations",
        "Week 4 evaluation handoff",
    ):
        add_bullet(document, item)

    add_heading(document, "Project overview")
    document.add_paragraph(
        "The Code Review Agent accepts one canonical GitHub pull-request URL, records the exact "
        "head SHA, and constructs a bounded review context from metadata, changed files, patches, "
        "and selected head-side content. LangGraph runs correctness, security, and test-coverage "
        "specialists concurrently. Deterministic Python code then validates, grounds, filters, "
        "deduplicates, and orders the findings."
    )
    document.add_paragraph(
        "The user approves potentially paid model calls before analysis and approves the exact "
        "Markdown body before publication. A separate publisher revalidates the head SHA and "
        "creates or updates one marker-tagged comment. Publishing is disabled by default, and the "
        "graph cannot access the write component."
    )
    add_heading(document, "Project success measure", 2)
    document.add_paragraph(
        "An engineer can inspect the captured pull-request state, see each specialist's status, "
        "review one deterministic grounded draft, and deliberately approve or reject that exact "
        "body without risking an accidental GitHub write."
    )

    document.add_page_break()
    add_heading(document, "Agent workflow")
    steps = (
        "Validate the canonical GitHub URL and preview metadata.",
        "Retrieve bounded changed-file data and capture the exact head SHA.",
        "Redact likely credentials and fence repository content as untrusted data.",
        "Run nonexecuting Python AST checks when parseable content is available.",
        "Fan out to correctness, security, and test-coverage specialists.",
        "Apply deterministic grounding, relevance, quality, contradiction, and duplicate gates.",
        "Render one marker-tagged Markdown review and show specialist and tool status.",
        "Bind human approval to URL, SHA, generated draft, and edited body.",
        "Revalidate the SHA and update or create one same-author marker comment when enabled.",
    )
    for number, step in enumerate(steps, start=1):
        add_number(document, number, step)

    add_heading(document, "Human handoffs", 2)
    add_table(
        document,
        ["Point", "User decision", "Effect"],
        [
            ["Before analysis", "Approve the displayed model call bound", "Enables Run review"],
            [
                "Before publication",
                "Approve the exact Markdown body",
                "Creates an approval fingerprint",
            ],
            ["On any change", "Review the new URL, SHA, or body", "Clears the previous approval"],
        ],
        [1.25, 2.75, 2.75],
    )

    add_heading(document, "System architecture")
    diagram = document.add_paragraph()
    diagram.alignment = WD_ALIGN_PARAGRAPH.CENTER
    diagram.paragraph_format.space_before = Pt(5)
    diagram.paragraph_format.space_after = Pt(10)
    run = diagram.add_run(
        "GitHub reads  →  bounded context  →  LangGraph specialists\n"
        "                                      ↓\n"
        "approval  ←  deterministic Markdown  ←  validation\n"
        "    ↓\n"
        "separate publisher  →  SHA check  →  one marker comment"
    )
    run.font.name = "Aptos Mono"
    run.font.size = Pt(9.5)

    add_heading(document, "Read and context components", 2)
    document.add_paragraph(
        "GitHubReadClient exposes GET operations only. It paginates changed files, retries bounded "
        "transient failures, and returns typed metadata and content. Context construction enforces "
        "file, patch, content, and total-character limits. Removed files, patchless files, and "
        "truncated patches cannot supply publishable head-side coordinates."
    )
    add_heading(document, "Model and graph components", 2)
    document.add_paragraph(
        "StructuredModelClient sends one capability probe using the exact schema used for reviews. "
        "The provider schema uses a conservative keyword set. Pydantic performs stricter internal "
        "validation. The graph has three independent nodes and one finalizer; it has no publisher, "
        "write token, or GitHub write tool."
    )
    add_heading(document, "Deterministic review components", 2)
    document.add_paragraph(
        "A finding must name a changed file, use a valid head-side range, overlap a changed line, "
        "name a subject present in the referenced hunk, include subject-specific evidence, and "
        "give a concrete recommendation. Direct changed tests can suppress matching test gaps. "
        "Stable sorting and duplicate precedence produce byte-identical Markdown."
    )

    add_heading(document, "Specialist contracts")
    add_table(
        document,
        ["Specialist", "Primary focus", "Failure effect"],
        [
            [
                "Correctness",
                "Concrete behavior defects and broken invariants",
                "One failure yields a partial review",
            ],
            [
                "Security",
                "Trust boundaries and exploitable unsafe behavior",
                "Two total failures abort publication",
            ],
            [
                "Test coverage",
                "Important changed behavior without a direct test",
                "Routine suggestions are suppressed",
            ],
        ],
        [1.3, 3.1, 2.35],
    )
    document.add_paragraph(
        "Each specialist returns at most three findings and three short validation notes. "
        "Category, severity, reviewer identity, string size, path, and coordinate constraints are "
        "validated. "
        "A test gap can be medium, low, or informational, but cannot be high solely because the "
        "hypothetical defect would be severe."
    )

    add_heading(document, "Security and human control")
    add_table(
        document,
        ["Risk", "Control"],
        [
            [
                "Prompt injection",
                "Untrusted-data fencing, no model tools, deterministic path and line checks",
            ],
            [
                "Credential exposure",
                "Pre-model pattern redaction and source-free logs and evaluation artifacts",
            ],
            [
                "Repository code execution",
                "No clone, import, package install, hook, test, or build; AST parses text only",
            ],
            ["Stale approval", "Fingerprint includes URL, SHA, generated body, and edited body"],
            [
                "Accidental write",
                "Default-off flag, separate token, controller gate, per-mutation SHA recheck",
            ],
            [
                "Duplicate comments",
                "Stable hidden marker, same-author update, timeout search before retry",
            ],
        ],
        [2.0, 4.75],
    )
    document.add_paragraph(
        "The publisher never overwrites a foreign-authored marker and stops when it finds "
        "multiple same-author marker matches. The evaluator constructs only read and model "
        "components and is "
        "incapable of publishing."
    )

    add_heading(document, "Testing and evaluation")
    add_heading(document, "Offline regression evidence", 2)
    document.add_paragraph(
        "The offline suite covers URL parsing, GitHub retries and pagination, context bounds "
        "and redaction, prompt-injection fencing, provider output, malformed retries, partial "
        "and abort "
        "paths, severity rules, grounding, relevance, direct-test contradiction, deduplication, "
        "deterministic Markdown, approval resets, publication behavior, evaluation artifacts, and "
        "the Streamlit workflow. HTTP and model boundaries use injected responses; tests make no "
        "paid calls."
    )
    document.add_paragraph(
        "Final verification passed 175 tests with 92.60 percent branch coverage, plus Ruff lint "
        "and formatting, two secret scans, the dependency audit, package and Twine checks, clean "
        "artifact installs, and source, wheel, and source-distribution Streamlit health requests."
    )
    add_heading(document, "Black box sequence", 2)
    for number, item in enumerate(
        (
            "Run offline tests and security checks.",
            "Preflight all ten pinned head SHAs before the first model call.",
            "Run both clean controls and stop on any finding or operational failure.",
            "Run all seven seeded defects and require one-to-one expected matches.",
            "Run the one controlled malformed-recovery case and require a partial review with "
            "exactly two completed specialists.",
            "Combine one result per case without rerunning for a better score.",
        ),
        start=1,
    ):
        add_number(document, number, item)
    document.add_paragraph(
        "The clean controls run before the seeded defects, and malformed recovery runs last. Every "
        "case runs at most once. A failed whole-catalog preflight makes zero model calls. Metrics "
        "include recall, precision, severity and location agreement, duplicate rate, specialist "
        "completion, partial recovery, and clean behavior. Unavailable metrics are N/A."
    )
    add_heading(document, "Golden dataset", 2)
    add_table(
        document,
        ["Case", "Role", "Pull request", "Expected result"],
        [
            ["01", "Seeded", "PR 1", "Correctness, security, and test gap"],
            ["02", "Clean", "PR 2", "No publishable finding"],
            ["03", "Seeded", "PR 3", "Correctness"],
            ["04", "Seeded", "PR 6", "Correctness"],
            ["05", "Seeded", "PR 4", "Security"],
            ["06", "Seeded", "PR 5", "Security"],
            ["07", "Seeded", "PR 7", "Test gap"],
            ["08", "Seeded", "PR 9", "Correctness"],
            ["09", "Clean", "PR 8", "No publishable finding"],
            ["10", "Recovery", "PR 10", "Correctness with one specialist failure"],
        ],
        [0.55, 0.9, 1.0, 4.3],
        font_size=8.5,
        vertical_margin=45,
    )
    document.add_paragraph(
        "The version 2 catalog contains seven seeded defects, two clean controls, and one "
        "controlled malformed-recovery case. All ten pull requests are open, pinned to unique "
        "head SHAs, and use neutral metadata. Fixture creation did not call the model, publish a "
        "review comment, or merge a fixture."
    )
    document.add_paragraph(
        "The authorized sequence preflighted all ten pins and stopped with "
        "clean-controls-produced-findings after golden-02 produced one high-severity correctness "
        "false positive. Golden-09 remained clean. The remaining seeded and recovery cases were "
        "then run once in diagnostic modes without overriding the failed official gate."
    )
    add_heading(document, "Fixture links", 3)
    fixture_links = (
        "01  https://github.com/melepatt/code-review-agent-demo/pull/1",
        "02  https://github.com/melepatt/code-review-agent-demo/pull/2",
        "03  https://github.com/melepatt/code-review-agent-demo/pull/3",
        "04  https://github.com/melepatt/code-review-agent-demo/pull/6",
        "05  https://github.com/melepatt/code-review-agent-demo/pull/4",
        "06  https://github.com/melepatt/code-review-agent-demo/pull/5",
        "07  https://github.com/melepatt/code-review-agent-demo/pull/7",
        "08  https://github.com/melepatt/code-review-agent-demo/pull/9",
        "09  https://github.com/melepatt/code-review-agent-demo/pull/8",
        "10  https://github.com/melepatt/code-review-agent-demo/pull/10",
    )
    for fixture_link in fixture_links:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        link_run = paragraph.add_run(fixture_link)
        link_run.font.name = "Aptos Mono"
        link_run.font.size = Pt(8.5)

    add_heading(document, "Evidence status")
    add_table(
        document,
        ["Evidence class", "Status", "Interpretation"],
        [
            [
                "Offline regression",
                "Passed",
                "175 tests and 92.60 percent branch coverage",
            ],
            [
                "Golden fixture setup",
                "Verified",
                "Ten open pinned PRs; no model call, comment, or merge",
            ],
            [
                "Official live evaluation",
                "Failed",
                "Stopped at clean controls: clean-controls-produced-findings",
            ],
            [
                "Diagnostic live coverage",
                "Recorded",
                "Ten unique cases across clean, seeded, and recovery artifacts",
            ],
            [
                "Controlled GitHub publication",
                "Not run",
                "No external comment was created or updated",
            ],
        ],
        [1.75, 1.15, 3.85],
    )
    document.add_paragraph(
        "The diagnostic aggregate does not convert the stopped sequence into a pass. GitHub review "
        "publication: Not run. A live model run does not authorize an external write."
    )
    add_heading(document, "Live evaluation baseline", 2)
    add_table(
        document,
        ["Measure", "Result", "Interpretation"],
        [
            ["Recall", "10 percent", "One of ten expected findings matched"],
            ["Precision", "14.29 percent", "One of seven accepted findings matched"],
            ["Specialist completion", "93.33 percent", "28 of 30 specialist slots completed"],
            ["Duplicate rate", "0 percent", "No duplicate accepted coordinates"],
            ["Clean behavior", "Failed", "Golden-02 produced one false positive"],
            ["Partial recovery", "Failed", "Golden-10 completed 1/3 specialists"],
            ["Severity agreement", "100 percent", "Based on one matched finding only"],
            ["Location agreement", "100 percent", "Based on one matched finding only"],
        ],
        [1.65, 1.25, 3.85],
        font_size=8.5,
        vertical_margin=70,
    )

    add_heading(document, "Development iterations")
    document.add_paragraph(
        "A builder review identified seven decisions that affected safety and tests. The design "
        "now uses head-side coordinates, rejects deleted-only evidence, includes the marker in "
        "the body "
        "shown for approval, updates only same-author comments, retains diagnostics without a body "
        "after two failures, and validates configuration per capability."
    )
    document.add_paragraph(
        "Test-first development found two implementation defects. Hunk relevance initially looked "
        "only at the reported lines and rejected a subject declared earlier in the same hunk. The "
        "validator now evaluates the complete referenced hunk. Duplicate rate initially compared "
        "the same coordinates across different pull requests; metric keys now include the case ID."
    )
    document.add_paragraph(
        "The Streamlit design added a model-call approval checkbox after the initial interface "
        "plan. "
        "The control states one capability probe and the configured retry bound before any paid "
        "operation can start."
    )
    document.add_paragraph(
        "A final code review prompted focused regression fixes: snapshot collection now rejects "
        "head drift, every remote mutation rechecks the SHA, capability probing is serialized, "
        "evaluation matches expected findings one-to-one, prompt fences survive truncation, and "
        "the wheel and source distribution include the page, launcher, catalog, and configuration."
    )

    document.add_page_break()
    add_heading(document, "Operation")
    add_heading(document, "Configuration", 2)
    add_table(
        document,
        ["Setting", "Required for", "Safe default"],
        [
            ["LLM_API_KEY and LLM_BASE_URL", "Model review", "Review disabled"],
            ["LLM_MODEL", "Model selection", "Qwen Qwen3 30B A3B Instruct 2507"],
            ["GITHUB_READ_TOKEN", "Higher GitHub read limits", "Anonymous public reads"],
            ["ENABLE_GITHUB_PUBLISH", "Any GitHub write", "false"],
            ["GITHUB_WRITE_TOKEN", "Comment create or update", "Absent"],
        ],
        [2.1, 2.45, 2.2],
    )
    add_heading(document, "Run commands", 2)
    commands = (
        "uv venv --python 3.12",
        "uv pip install --python .venv/bin/python -e '.[dev]'",
        "cp .env.example .env",
        ".venv/bin/streamlit run streamlit_app.py",
        ".venv/bin/python scripts/run_evaluation.py --mode validate",
        ".venv/bin/python scripts/run_evaluation.py --mode sequence --approve-model-calls",
    )
    for command in commands:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(command)
        run.font.name = "Aptos Mono"
        run.font.size = Pt(9.5)

    document.add_page_break()
    add_heading(document, "Results and limitations")
    document.add_paragraph(
        "Offline evidence supports the implemented control flow and safety boundaries. The live "
        "evaluation separately shows that the configured model did not meet the golden gates. The "
        "official sequence failed on a clean false positive. Diagnostic runs covered the remaining "
        "cases once, producing 10 percent recall, 14.29 percent precision, 93.33 percent "
        "specialist completion, failed clean behavior, and failed recovery. Severity and location "
        "agreement "
        "are 100 percent for one matched finding only."
    )
    add_bullet(
        document, "Bounded context can omit repository-wide behavior needed to confirm a finding."
    )
    add_bullet(document, "Credential patterns cannot recognize every secret format.")
    add_bullet(
        document,
        "Conservative gates can suppress a valid insight when patch evidence is incomplete.",
    )
    add_bullet(
        document, "Ten synthetic evaluation cases cannot support a production-accuracy claim."
    )
    add_bullet(
        document,
        "Source-free artifacts do not show whether every miss began in generation, structured "
        "output, or deterministic rejection.",
    )
    add_bullet(
        document, "Operators must use narrowly scoped GitHub tokens and a controlled target PR."
    )

    add_heading(document, "Week 4 evaluation handoff")
    document.add_paragraph(
        "Week 4 should preserve this run as the baseline rather than tune against it informally. "
        "The next evaluation iteration should audit expected labels and acceptable secondary "
        "findings, add source-free counts for generated and rejected findings and retry failures, "
        "and expand the dataset with balanced clean and seeded cases."
    )
    document.add_paragraph(
        "Prompt and model experiments should use a development split. A held-out split should "
        "measure final accuracy, and predetermined repeated runs should report variation without "
        "rerunning individual cases for a better score."
    )

    add_heading(document, "Demonstration plan")
    document.add_paragraph(
        "The demonstration previews a public pull request, shows the exact SHA and model-call "
        "approval boundary, and uses an already captured result when available to explain "
        "specialist status, deterministic validation, exact-body approval, and approval reset. It "
        "then shows the saved clean, seeded, and recovery artifacts, explains the failed official "
        "gate and diagnostic baseline, and closes with the Week 4 handoff. Golden cases are not "
        "rerun solely for recording, and GitHub publication remains disabled."
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("usage: build_project_docx.py OUTPUT.docx")
    build(Path(sys.argv[1]))
